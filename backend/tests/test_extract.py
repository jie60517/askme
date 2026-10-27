# -*- coding: utf-8 -*-
"""上傳檔轉純文字(extract.extract_text)。

這裡守的是一種**靜默**的錯:抽不到的內容不會報錯,claude 只會看到比較少的東西,
然後照樣給一個很肯定的答案。所以每一條都在斷言「那段字真的在裡面」。
"""
import struct

import pytest

from app import extract


# ---------------------------------------------------------------- Word

def test_docx_的表格不可以被丟掉(tmp_path):
    """python-docx 的 `doc.paragraphs` **不含表格**,而很多文件的內容幾乎全在表格裡。"""
    from docx import Document

    doc = Document()
    doc.add_paragraph("前言")
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "欄位代號"
    table.rows[0].cells[1].text = "item_code"
    table.rows[1].cells[0].text = "型態"
    table.rows[1].cells[1].text = "C002"
    doc.add_paragraph("結語")
    path = tmp_path / "fields.docx"
    doc.save(str(path))

    text = extract.extract_text(path)

    assert "item_code" in text and "C002" in text
    # 順序要跟文件一致,不然「這張表在講哪一段」就斷了
    assert text.index("前言") < text.index("item_code") < text.index("結語")


# ---------------------------------------------------------------- PowerPoint

def test_pptx_的表格與備忘稿不可以被丟掉(tmp_path):
    """表格圖形沒有 text_frame,只看 has_text_frame 的話那幾頁是整頁空白。"""
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "欄位一覽"
    table = slide.shapes.add_table(1, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
    table.cell(0, 0).text = "item_code"
    table.cell(0, 1).text = "品號"
    slide.notes_slide.notes_text_frame.text = "這頁要講的重點"
    path = tmp_path / "deck.pptx"
    prs.save(str(path))

    text = extract.extract_text(path)

    assert "欄位一覽" in text
    assert "item_code" in text and "品號" in text
    assert "這頁要講的重點" in text


def test_pptx_的群組圖形要展開(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    boxes = []
    for i in range(2):
        box = slide.shapes.add_textbox(Inches(1 + i), Inches(1), Inches(1), Inches(1))
        box.text_frame.text = f"群組裡的字{i}"
        boxes.append(box)
    slide.shapes.add_group_shape(boxes)
    path = tmp_path / "group.pptx"
    prs.save(str(path))

    text = extract.extract_text(path)

    assert "群組裡的字0" in text and "群組裡的字1" in text


# ---------------------------------------------------------------- 舊版 Excel

def _biff2_xls(rows):
    """手工組一份最小的 .xls(BIFF2)。

    xlrd 只讀不寫,而 xlwt 不是這個專案的相依 —— 為了測一個讀取路徑多裝一支
    寫檔函式庫不划算。BIFF2 的記錄格式很單純:BOF、CODEPAGE、每格一筆、EOF。
    """
    def rec(op, data):
        return struct.pack("<HH", op, len(data)) + data

    out = rec(0x0009, struct.pack("<HH", 2, 0x0010)) + rec(0x0042, struct.pack("<H", 950))
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            head = struct.pack("<HH", r, c) + b"\x00\x00\x00"
            if isinstance(value, (int, float)):
                out += rec(0x0003, head + struct.pack("<d", value))
            else:
                raw = value.encode("cp950")
                out += rec(0x0004, head + bytes([len(raw)]) + raw)
    return out + rec(0x000A, b"")


def test_舊版_xls_讀得出內容(tmp_path):
    path = tmp_path / "報表.xls"
    path.write_bytes(_biff2_xls([["品號", "數量"], ["A-100", 3]]))

    text = extract.extract_text(path)

    assert "品號\t數量" in text
    # 整數不要變成 "3.0" —— xlrd 的數字一律是 float
    assert "A-100\t3" in text


def test_讀不了的舊版_xls_給說明而不是亂碼(tmp_path):
    path = tmp_path / "報表.xls"
    path.write_bytes(extract._OLE2_MAGIC + b"\x00" * 200)

    text = extract.extract_text(path)

    assert "Excel 97-2003" in text and ".xlsx" in text


def test_xlrd_沒裝的時候降級成說明(tmp_path, monkeypatch):
    """部署環境的 venv 沒跟上 requirements 時:功能不見了,但使用者知道該怎麼辦。"""
    def no_xlrd(path):
        raise ImportError("No module named 'xlrd'")

    monkeypatch.setattr(extract, "_extract_xls", no_xlrd)
    path = tmp_path / "報表.xls"
    path.write_bytes(_biff2_xls([["品號"]]))
    # BIFF2 不是 OLE2,所以要模擬真的舊檔才會走到說明那條
    path.write_bytes(extract._OLE2_MAGIC + path.read_bytes())

    assert "Excel 97-2003" in extract.extract_text(path)


# ---------------------------------------------------------------- 舊格式的說明

@pytest.mark.parametrize(
    "name,expected",
    [("報告.doc", "Word 97-2003"), ("簡報.ppt", "PowerPoint 97-2003")],
)
def test_舊格式的二進位檔要明講讀不到(tmp_path, name, expected):
    """不講的話 claude 會去讀原始檔,拿到亂碼再掰一個答案出來。"""
    path = tmp_path / name
    path.write_bytes(extract._OLE2_MAGIC + b"\x00" * 200)

    text = extract.extract_text(path)

    assert expected in text
    assert "不要猜測" in text


def test_副檔名騙人的匯出檔不要被當成舊格式(tmp_path):
    """很多系統匯出的「.xls」其實是 tab 分隔的純文字 —— 那種讓 claude 直接讀就對了。"""
    path = tmp_path / "匯出.xls"
    path.write_text("品號\t數量\nA-100\t3\n", encoding="cp950")

    assert extract.extract_text(path) is None


def test_不支援的格式回_None(tmp_path):
    path = tmp_path / "a.zip"
    path.write_bytes(b"PK\x03\x04rest")

    assert extract.extract_text(path) is None
