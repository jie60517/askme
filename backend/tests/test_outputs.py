# -*- coding: utf-8 -*-
"""收尾階段的產檔:PDF、簡報、打包,以及子資料夾攤平。

共同的三條規矩(每一支都要守):
  1. 產出來的檔案要**留在第一層**,不然 list_output_files 掃不到 = 下載清單裡沒有它。
  2. manifest 用完就刪,不要出現在使用者的下載清單裡。
  3. 失敗只回訊息、不丟例外 —— 收尾時答案已經算完了。
"""
import json
import zipfile
from pathlib import Path

import pytest

from app import agent, config, docgen, pdfgen, slidegen, zipgen


@pytest.fixture
def out_dir(tmp_path, monkeypatch):
    output = tmp_path / "outputs"
    output.mkdir()
    monkeypatch.setattr(config, "OUTPUT_DIR", output)
    session = output / "s-1"
    session.mkdir()
    return session


# ---------------------------------------------------------------- PDF

MARKDOWN = """# 出貨流程

第一段說明,含 **粗體** 與 `checkout.py`。

## 欄位

| 欄位 | 說明 |
|---|---|
| item_code | 品號 |

- 第一點
- 第二點

```
print(order_total())
```
"""


def test_md_轉成_pdf(out_dir):
    (out_dir / "流程.md").write_text(MARKDOWN, encoding="utf-8")
    (out_dir / pdfgen.MANIFEST_NAME).write_text("流程.md\n", encoding="utf-8")

    assert pdfgen.apply_pdf_manifest("s-1") == []

    pdf = out_dir / "流程.pdf"
    assert pdf.is_file() and pdf.read_bytes().startswith(b"%PDF")
    assert not (out_dir / pdfgen.MANIFEST_NAME).exists()  # 清單不留在下載清單裡
    assert (out_dir / "流程.md").is_file()  # 原檔留著


def test_pdf_清單裡的路徑只認檔名(out_dir):
    """讀別的地方的檔案是 copy manifest 的事,那邊才有來源白名單。"""
    (out_dir / pdfgen.MANIFEST_NAME).write_text("../../../etc/passwd\n", encoding="utf-8")

    notices = pdfgen.apply_pdf_manifest("s-1")

    assert len(notices) == 1 and "沒有轉成 PDF" in notices[0]
    assert not list(out_dir.glob("*.pdf"))


def test_pdf_產出失敗不丟例外(out_dir, monkeypatch):
    def boom(markdown, out_path, title=None):
        out_path.write_bytes(b"half")
        raise RuntimeError("壞了")

    monkeypatch.setattr(pdfgen, "render", boom)
    (out_dir / "a.md").write_text("x", encoding="utf-8")
    (out_dir / pdfgen.MANIFEST_NAME).write_text("a.md", encoding="utf-8")

    notices = pdfgen.apply_pdf_manifest("s-1")

    assert len(notices) == 1 and "轉成 PDF 失敗" in notices[0]
    assert not (out_dir / "a.pdf").exists()  # 半成品要清掉


def test_沒有清單就什麼都不做(out_dir):
    assert pdfgen.apply_pdf_manifest("s-1") == []


def test_表格長短不一也產得出來(out_dir):
    """模型寫的表格少一根 | 是常有的事,reportlab 的 Table 不接受長短不一的列。"""
    (out_dir / "a.md").write_text(
        "| a | b | c |\n|---|---|---|\n| 1 | 2 |\n", encoding="utf-8"
    )
    (out_dir / pdfgen.MANIFEST_NAME).write_text("a.md", encoding="utf-8")

    assert pdfgen.apply_pdf_manifest("s-1") == []
    assert (out_dir / "a.pdf").is_file()


# ---------------------------------------------------------------- 簡報

def test_slides_轉成_pptx(out_dir):
    spec = {
        "title": "出貨流程說明",
        "subtitle": "2026 版",
        "slides": [
            {"heading": "流程", "bullets": ["下單", "出貨"], "notes": "這頁講重點"},
            {
                "heading": "欄位",
                "table": {"header": ["欄位", "說明"], "rows": [["item_code", "品號"]]},
            },
        ],
    }
    (out_dir / slidegen.MANIFEST_NAME).write_text(
        json.dumps(spec, ensure_ascii=False), encoding="utf-8"
    )

    assert slidegen.apply_slides_manifest("s-1") == []

    from pptx import Presentation

    prs = Presentation(str(out_dir / "出貨流程說明.pptx"))
    assert len(prs.slides) == 3  # 封面 + 兩頁
    texts = [
        shape.text_frame.text
        for slide in prs.slides
        for shape in slide.shapes
        if shape.has_text_frame
    ]
    assert any("出貨流程說明" in t for t in texts)
    assert any("下單" in t for t in texts)
    assert prs.slides[1].notes_slide.notes_text_frame.text == "這頁講重點"
    # 表格那頁
    assert any(shape.has_table for shape in prs.slides[2].shapes)
    assert not (out_dir / slidegen.MANIFEST_NAME).exists()


def test_slides_資料壞掉只回訊息(out_dir):
    (out_dir / slidegen.MANIFEST_NAME).write_text("{壞掉的 json", encoding="utf-8")

    notices = slidegen.apply_slides_manifest("s-1")

    assert len(notices) == 1 and "不是合法的 JSON" in notices[0]
    assert not list(out_dir.glob("*.pptx"))


def test_slides_一頁太多條列會截斷並說明(out_dir):
    spec = {"title": "t", "slides": [{"heading": "多", "bullets": [str(i) for i in range(30)]}]}
    (out_dir / slidegen.MANIFEST_NAME).write_text(
        json.dumps(spec, ensure_ascii=False), encoding="utf-8"
    )

    notices = slidegen.apply_slides_manifest("s-1")

    assert len(notices) == 1 and "條列超過" in notices[0]
    assert (out_dir / "t.pptx").is_file()  # 還是要產得出來


# ---------------------------------------------------------------- 打包

def test_打包成_zip(out_dir):
    (out_dir / "a.csv").write_text("x", encoding="utf-8")
    (out_dir / "b.md").write_text("y", encoding="utf-8")
    (out_dir / zipgen.MANIFEST_NAME).write_text("交付檔案\na.csv\nb.md\n", encoding="utf-8")

    assert zipgen.apply_zip_manifest("s-1") == []

    with zipfile.ZipFile(out_dir / "交付檔案.zip") as zf:
        assert sorted(zf.namelist()) == ["a.csv", "b.md"]
    assert not (out_dir / zipgen.MANIFEST_NAME).exists()


def test_打包只收這一題資料夾裡的檔案(out_dir, tmp_path):
    outside = tmp_path / "secret.txt"
    outside.write_text("秘密", encoding="utf-8")
    (out_dir / "a.csv").write_text("x", encoding="utf-8")
    (out_dir / zipgen.MANIFEST_NAME).write_text(
        "包.zip\na.csv\n%s\n" % outside, encoding="utf-8"
    )

    notices = zipgen.apply_zip_manifest("s-1")

    assert len(notices) == 1 and "不是這一題產出的檔案" in notices[0]
    with zipfile.ZipFile(out_dir / "包.zip") as zf:
        assert zf.namelist() == ["a.csv"]


def test_打包清單不完整就不打包(out_dir):
    (out_dir / zipgen.MANIFEST_NAME).write_text("只有檔名.zip\n", encoding="utf-8")

    notices = zipgen.apply_zip_manifest("s-1")

    assert len(notices) == 1 and "不完整" in notices[0]
    assert not list(out_dir.glob("*.zip"))


def test_打包太大就不壓(out_dir, monkeypatch):
    monkeypatch.setattr(zipgen, "MAX_TOTAL_BYTES", 10)
    (out_dir / "a.csv").write_text("x" * 100, encoding="utf-8")
    (out_dir / zipgen.MANIFEST_NAME).write_text("包\na.csv", encoding="utf-8")

    notices = zipgen.apply_zip_manifest("s-1")

    assert len(notices) == 1 and "超過" in notices[0]
    assert not list(out_dir.glob("*.zip"))


# ---------------------------------------------------------------- 子資料夾

def test_子資料夾裡的檔案要搬到第一層(out_dir):
    """不搬的話 list_output_files 掃不到 —— 檔案寫成功了,使用者卻什麼都看不到。"""
    nested = out_dir / "報表" / "細項"
    nested.mkdir(parents=True)
    (nested / "a.csv").write_text("x", encoding="utf-8")

    assert agent.flatten_output_dir("s-1") == []

    assert (out_dir / "a.csv").is_file()
    assert not (out_dir / "報表").exists()  # 空目錄收掉
    assert agent.list_output_files("s-1") == ["a.csv"]


def test_搬上來撞名就加序號(out_dir):
    (out_dir / "a.csv").write_text("第一層", encoding="utf-8")
    (out_dir / "子").mkdir()
    (out_dir / "子" / "a.csv").write_text("子資料夾", encoding="utf-8")

    agent.flatten_output_dir("s-1")

    assert (out_dir / "a.csv").read_text(encoding="utf-8") == "第一層"
    assert (out_dir / "a(2).csv").read_text(encoding="utf-8") == "子資料夾"


def test_沒有子資料夾就什麼都不動(out_dir):
    (out_dir / "a.csv").write_text("x", encoding="utf-8")

    assert agent.flatten_output_dir("s-1") == []
    assert agent.list_output_files("s-1") == ["a.csv"]


def test_slides_條列太長時表格不硬塞(out_dir):
    """硬塞的話表格會長到頁面外面,而那在產檔當下完全看不出來。"""
    spec = {
        "title": "t",
        "slides": [
            {
                "heading": "滿了",
                "bullets": [f"第 {i} 點" for i in range(9)],
                "table": {"header": ["a", "b"], "rows": [["1", "2"], ["3", "4"]]},
            }
        ],
    }
    (out_dir / slidegen.MANIFEST_NAME).write_text(
        json.dumps(spec, ensure_ascii=False), encoding="utf-8"
    )

    notices = slidegen.apply_slides_manifest("s-1")

    assert any("排不下" in n for n in notices)
    from pptx import Presentation

    prs = Presentation(str(out_dir / "t.pptx"))
    assert not any(shape.has_table for shape in prs.slides[1].shapes)


# ---------------------------------------------------------------- Word

def _docx_text(path: Path) -> str:
    from docx import Document

    doc = Document(str(path))
    cells = [c.text for t in doc.tables for row in t.rows for c in row.cells]
    return "\n".join([p.text for p in doc.paragraphs] + cells)


def test_md_轉成_docx(out_dir):
    (out_dir / "流程.md").write_text(MARKDOWN, encoding="utf-8")
    (out_dir / docgen.MANIFEST_NAME).write_text("流程.md\n", encoding="utf-8")

    assert docgen.apply_docx_manifest("s-1") == []

    docx = out_dir / "流程.docx"
    assert docx.is_file() and docx.read_bytes().startswith(b"PK")  # docx 是 zip
    text = _docx_text(docx)
    for expected in ("出貨流程", "粗體", "checkout.py", "item_code", "品號", "第二點", "print(order_total())"):
        assert expected in text
    assert "**" not in text  # 行內標記要變成格式,不是原字
    assert not (out_dir / docgen.MANIFEST_NAME).exists()  # 清單不留在下載清單裡
    assert (out_dir / "流程.md").is_file()  # 原檔留著


def test_docx_中文字型寫在_eastAsia(out_dir):
    """只設 font.name 的話中文會用 Word 預設字型;標題樣式還綁著佈景主題字型。"""
    from docx import Document
    from docx.oxml.ns import qn

    (out_dir / "a.md").write_text("# 標題\n\n內文\n", encoding="utf-8")
    (out_dir / docgen.MANIFEST_NAME).write_text("a.md", encoding="utf-8")
    docgen.apply_docx_manifest("s-1")

    doc = Document(str(out_dir / "a.docx"))
    for name in ("Normal", "Heading 1"):
        rfonts = doc.styles[name].element.rPr.rFonts
        assert rfonts.get(qn("w:eastAsia")) == docgen.CJK_FONT
        assert rfonts.get(qn("w:asciiTheme")) is None


def test_docx_引用不留下大於符號(out_dir):
    """模型常用 `> 適用範圍說明` 開頭;照原字印的話 Word 裡會留一個 `>`。"""
    from docx import Document

    (out_dir / "a.md").write_text("> 第一行\n> **第二行**\n\n內文\n", encoding="utf-8")
    (out_dir / docgen.MANIFEST_NAME).write_text("a.md", encoding="utf-8")
    docgen.apply_docx_manifest("s-1")

    quotes = [p for p in Document(str(out_dir / "a.docx")).paragraphs if p.style.name == "Quote"]
    assert len(quotes) == 1 and quotes[0].text == "第一行 第二行"
    assert ">" not in _docx_text(out_dir / "a.docx")


def test_docx_表格長短不一也產得出來(out_dir):
    (out_dir / "a.md").write_text("| a | b | c |\n|---|---|---|\n| 1 | 2 |\n", encoding="utf-8")
    (out_dir / docgen.MANIFEST_NAME).write_text("a.md", encoding="utf-8")

    assert docgen.apply_docx_manifest("s-1") == []
    assert "2" in _docx_text(out_dir / "a.docx")


def test_docx_清單裡的路徑只認檔名(out_dir):
    (out_dir / docgen.MANIFEST_NAME).write_text("../../../etc/passwd\n", encoding="utf-8")

    notices = docgen.apply_docx_manifest("s-1")

    assert len(notices) == 1 and "沒有轉成 Word" in notices[0]
    assert not list(out_dir.glob("*.docx"))


def test_docx_產出失敗不丟例外(out_dir, monkeypatch):
    def boom(markdown, out_path, title=None):
        out_path.write_bytes(b"half")
        raise RuntimeError("壞了")

    monkeypatch.setattr(docgen, "render", boom)
    (out_dir / "a.md").write_text("x", encoding="utf-8")
    (out_dir / docgen.MANIFEST_NAME).write_text("a.md", encoding="utf-8")

    notices = docgen.apply_docx_manifest("s-1")

    assert len(notices) == 1 and "轉成 Word 失敗" in notices[0]
    assert not (out_dir / "a.docx").exists()  # 半成品要清掉


def test_docx_不蓋掉同名的既有檔(out_dir):
    (out_dir / "手冊.docx").write_bytes(b"original")
    (out_dir / "手冊.md").write_text("# 手冊\n", encoding="utf-8")
    (out_dir / docgen.MANIFEST_NAME).write_text("手冊.md", encoding="utf-8")

    docgen.apply_docx_manifest("s-1")

    assert (out_dir / "手冊.docx").read_bytes() == b"original"
    assert len(list(out_dir.glob("*.docx"))) == 2


def test_同一份_md_可以同時轉_word_與_pdf(out_dir):
    (out_dir / "a.md").write_text("# 標題\n", encoding="utf-8")
    (out_dir / docgen.MANIFEST_NAME).write_text("a.md", encoding="utf-8")
    (out_dir / pdfgen.MANIFEST_NAME).write_text("a.md", encoding="utf-8")

    assert docgen.apply_docx_manifest("s-1") == []
    assert pdfgen.apply_pdf_manifest("s-1") == []
    assert (out_dir / "a.docx").is_file() and (out_dir / "a.pdf").is_file()


def test_提示裡有教_word_的做法():
    """沒講的話使用者要 Word 時,模型會回「沒辦法」或改給 .md / .pdf。"""
    prompt = agent.build_prompt("把答案整理成 Word", session_id="s-1", uploaded_filenames=None)
    assert docgen.MANIFEST_NAME in prompt and "Word" in prompt


# ---------------------------------------------------------------- 收尾的順序

def test_後處理一定要排在_list_output_files_前面():
    """這條規矩寫在五支產檔模組的 docstring 裡,這裡用測試守著它。

    排錯的症狀是「檔案產出來了,但下載清單裡沒有」—— 畫面上沒有錯誤,
    答案裡還說「已經產生」。這種錯沒有測試抓得到,只能釘原始碼的順序:
    收尾那支函式很短,而順序正是它唯一在做的事。
    """
    import inspect

    from app.routers import ask

    # 註解要先拿掉:那幾句「一定要排在 list_output_files 前面」的提醒本身就含關鍵字,
    # 不濾掉的話比到的是註解的位置而不是呼叫的位置。
    source = "\n".join(
        line.split("#")[0] for line in inspect.getsource(ask._finalize).splitlines()
    )
    listing = source.index("list_output_files")
    for call in (
        "apply_copy_manifest",
        "flatten_output_dir",
        "apply_csv_exports",
        "apply_docx_manifest",
        "apply_pdf_manifest",
        "apply_slides_manifest",
        "apply_zip_manifest",
    ):
        assert source.index(call) < listing, f"{call} 排到 list_output_files 後面了"
    # 打包要最後:前面幾支產出來的檔案也要包得進去
    for call in ("apply_csv_exports", "apply_docx_manifest", "apply_pdf_manifest", "apply_slides_manifest"):
        assert source.index(call) < source.index("apply_zip_manifest")
    # 攤平要排在 copy manifest **之後**(複製進來的檔案也要攤平)、其餘之前
    assert (
        source.index("apply_copy_manifest")
        < source.index("flatten_output_dir")
        < source.index("apply_csv_exports")
    )


# ---------------------------------------------------------------- 清單刪不掉

def test_清單刪不掉也不可以把整題弄成失敗(out_dir, monkeypatch):
    """unlink 自己會丟 OSError(檔案被鎖、權限沒了)。寫在 finally 裡的話會逃出去。"""
    def cannot_unlink(self, missing_ok=False):
        raise PermissionError(13, "Permission denied")

    (out_dir / "a.md").write_text("x", encoding="utf-8")
    (out_dir / pdfgen.MANIFEST_NAME).write_text("a.md", encoding="utf-8")
    monkeypatch.setattr(Path, "unlink", cannot_unlink)

    notices = pdfgen.apply_pdf_manifest("s-1")

    assert any("刪不掉" in n for n in notices)
    assert (out_dir / "a.pdf").is_file()  # 清單刪不掉不影響產檔


def test_清單讀不出來要講(out_dir, monkeypatch):
    """靜默的話,答案說「已經產生」而下載清單裡沒有,畫面上也沒有任何錯誤。"""
    def cannot_read(self, **kwargs):
        raise OSError(5, "I/O error")

    (out_dir / zipgen.MANIFEST_NAME).write_text("包.zip\na.csv", encoding="utf-8")
    monkeypatch.setattr(Path, "read_text", cannot_read)

    notices = zipgen.apply_zip_manifest("s-1")

    assert any("讀不出來" in n for n in notices)


def test_打包超過上限要講(out_dir, monkeypatch):
    monkeypatch.setattr(zipgen, "MAX_MEMBERS", 2)
    for i in range(4):
        (out_dir / f"a{i}.csv").write_text("x", encoding="utf-8")
    (out_dir / zipgen.MANIFEST_NAME).write_text(
        "包\n" + "\n".join(f"a{i}.csv" for i in range(4)), encoding="utf-8"
    )

    notices = zipgen.apply_zip_manifest("s-1")

    assert any("超過上限" in n for n in notices)
    with zipfile.ZipFile(out_dir / "包.zip") as zf:
        assert zf.namelist() == ["a0.csv", "a1.csv"]


def test_載不到中文字型要講(out_dir, monkeypatch):
    """不講的話使用者拿到一份中文整段空白的 PDF,而畫面上什麼提示都沒有。"""
    monkeypatch.setattr(pdfgen, "CJK_FONT", "這個字型不存在")
    (out_dir / "a.md").write_text("測試", encoding="utf-8")
    (out_dir / pdfgen.MANIFEST_NAME).write_text("a.md", encoding="utf-8")

    notices = pdfgen.apply_pdf_manifest("s-1")

    assert any("中文" in n for n in notices)
    assert (out_dir / "a.pdf").is_file()  # 檔案照產,英數的部分還讀得到


# ---------------------------------------------------------------- 撞名不覆蓋

def test_pdf_不可以蓋掉同名的既有檔(out_dir):
    """最常見的情況:copy manifest 剛把一份真的手冊.pdf 複製進來,模型又寫了手冊.md。

    蓋掉的話使用者要的正本就沒了,而且畫面上不會有任何跡象(回傳 [] 一句話都沒有)。
    """
    (out_dir / "手冊.pdf").write_bytes(b"%PDF-1.4 copy-manifest-original")
    (out_dir / "手冊.md").write_text("# 轉出來的", encoding="utf-8")
    (out_dir / pdfgen.MANIFEST_NAME).write_text("手冊.md", encoding="utf-8")

    pdfgen.apply_pdf_manifest("s-1")

    assert (out_dir / "手冊.pdf").read_bytes() == b"%PDF-1.4 copy-manifest-original"
    assert (out_dir / "手冊(2).pdf").read_bytes().startswith(b"%PDF")


def test_打包撞名時不要說成_不是這一題產出的檔案(out_dir):
    """去重要在挑成員之前做,不然清單裡列到同名的既有 .zip 會拿到誤導訊息。"""
    (out_dir / "包.zip").write_bytes("既有的壓縮檔".encode("utf-8"))
    (out_dir / "a.csv").write_text("x", encoding="utf-8")
    (out_dir / zipgen.MANIFEST_NAME).write_text("包.zip\na.csv\n包.zip\n", encoding="utf-8")

    notices = zipgen.apply_zip_manifest("s-1")

    assert not any("不是這一題產出的檔案" in n for n in notices)
    # 既有的不動
    assert (out_dir / "包.zip").read_bytes() == "既有的壓縮檔".encode("utf-8")
    with zipfile.ZipFile(out_dir / "包(2).zip") as zf:
        assert sorted(zf.namelist()) == ["a.csv", "包.zip"]


def test_slides_空的表格不要噴假訊息(out_dir):
    """使用者根本沒放表格,卻看到「表格排不下」只會一頭霧水。"""
    spec = {
        "title": "t",
        "slides": [{"heading": "只有條列", "bullets": [str(i) for i in range(9)], "table": {}}],
    }
    (out_dir / slidegen.MANIFEST_NAME).write_text(
        json.dumps(spec, ensure_ascii=False), encoding="utf-8"
    )

    notices = slidegen.apply_slides_manifest("s-1")

    assert not any("表格" in n for n in notices)
