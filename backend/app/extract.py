import datetime
import zipfile
from pathlib import Path
from typing import List, Optional

#: Office 檔(其實是 zip)解開之後所有成員加起來最多幾 bytes、單一成員最多幾 bytes。
#: deflate 可以壓到 1000:1 —— 一個 50 MB 的 .docx 裡面可以是 50 GB 的 XML,
#: python-docx / python-pptx 會把整份 XML 進 lxml。defusedxml 擋的是實體展開,
#: 擋不住這個。正常文件的 XML 本體幾 MB 就很大了;圖片是另外的成員,各自算。
_MAX_UNZIPPED_TOTAL = 256 * 1024 * 1024
_MAX_UNZIPPED_MEMBER = 128 * 1024 * 1024


def _zip_within_budget(path: Path) -> bool:
    """看 zip 目錄裡宣告的解壓大小,超過預算就不要解。**不會解壓任何東西。**

    `file_size` 是 central directory 裡寫的,讀它只要讀檔尾;壞掉的 zip 直接回 False
    (呼叫端會讓 claude 自己讀原始檔,跟對不支援格式的處理一樣)。
    """
    try:
        with zipfile.ZipFile(path) as zf:
            total = 0
            for info in zf.infolist():
                if info.file_size > _MAX_UNZIPPED_MEMBER:
                    return False
                total += info.file_size
                if total > _MAX_UNZIPPED_TOTAL:
                    return False
    except (zipfile.BadZipFile, OSError):
        return False
    return True


# OLE2 複合文件的檔頭(.doc/.xls/.ppt 97-2003 都是這個)。
#
# 用來分辨兩種長得一樣的東西:**真的**舊格式二進位檔,跟「副檔名叫 .xls、
# 內容其實是 HTML 或 tab 分隔文字」的系統匯出檔(這種很常見)。
# 前者讀原始檔只會拿到亂碼,要明講;後者讓 claude 直接讀反而完全正確。
_OLE2_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"

# 舊格式的說法。這會被當成「檔案內容」寫進 {檔名}.txt 交給 claude,
# 所以要寫成它能照著轉述給使用者的樣子 —— 不然它會拿亂碼硬掰一個答案出來。
_OLD_FORMAT_HINT = (
    "(系統訊息:這份檔案是舊版的 {name} 格式,後端讀不出它的內容,"
    "上面沒有任何檔案內容可以參考。請直接告訴使用者:"
    "請用 Office 開啟後另存成 {modern} 再重新上傳,或把需要的段落直接貼進對話。"
    "**不要猜測這份檔案裡有什麼,也不要嘗試自己去讀原始檔** —— 那是二進位檔,"
    "讀出來的是亂碼。)"
)

_OLD_FORMATS = {
    ".doc": ("Word 97-2003", ".docx"),
    ".xls": ("Excel 97-2003", ".xlsx"),
    ".ppt": ("PowerPoint 97-2003", ".pptx"),
}


def _is_ole2(path: Path) -> bool:
    try:
        with path.open("rb") as fh:
            return fh.read(8) == _OLE2_MAGIC
    except OSError:
        return False


def extract_text(path: Path) -> Optional[str]:
    """把常見的二進位文件格式先轉成純文字,讓 claude 不用自己想辦法解壓縮/解析。
    失敗或不支援的格式回傳 None,呼叫端會讓 claude 直接讀原始檔案。
    """
    suffix = path.suffix.lower()
    if suffix in (".pptx", ".docx", ".xlsx", ".xlsm") and not _zip_within_budget(path):
        return None
    try:
        if suffix == ".pptx":
            return _extract_pptx(path)
        if suffix == ".docx":
            return _extract_docx(path)
        if suffix in (".xlsx", ".xlsm"):
            return _extract_xlsx(path)
        if suffix == ".xls":
            # 舊版 Excel 仍然很常見。xlrd 讀得了它(2.x 也只讀得了它)。
            return _extract_xls(path)
    except Exception:
        # 這裡刻意不分「函式庫沒裝」與「檔案壞了」:兩種都往下走同一條路 ——
        # 真的是舊格式二進位就給說明,不是的話(副檔名騙人)還是讓 claude 讀原始檔。
        pass
    if suffix in _OLD_FORMATS and _is_ole2(path):
        name, modern = _OLD_FORMATS[suffix]
        return _OLD_FORMAT_HINT.format(name=name, modern=modern)
    return None


def _table_lines(rows: List[List[str]]) -> List[str]:
    """把一張表印成 tab 分隔的幾行。

    **表格一定要標出來。** 少了「== 表格 ==」這一行,欄位表跟內文黏成一片,
    claude 分不出「這是一列資料」還是「這是一句話」。
    """
    lines = ["== 表格 =="]
    for row in rows:
        cells = [" ".join(str(c or "").split()) for c in row]
        if any(cells):
            lines.append("\t".join(cells))
    return lines if len(lines) > 1 else []


def _extract_pptx(path: Path) -> str:
    """簡報轉文字:文字框、**表格**、**備忘稿**。

    ⚠️ **表格圖形沒有 `text_frame`**,只看 `has_text_frame` 的話「一頁只有一張表」
    的那幾頁會是**整頁空白** —— claude 看到一份沒有內容的簡報,卻照樣會給答案。
    """
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    def walk(shapes, parts: List[str]) -> None:
        for shape in shapes:
            # 群組圖形本身沒有文字,內容全在裡面 —— 不展開就整組不見
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                walk(shape.shapes, parts)
                continue
            if shape.has_table:
                rows = [[cell.text for cell in row.cells] for row in shape.table.rows]
                parts.extend(_table_lines(rows))
                continue
            if shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if text:
                    parts.append(text)

    prs = Presentation(str(path))
    parts: List[str] = []
    for i, slide in enumerate(prs.slides, 1):
        parts.append(f"== 第 {i} 頁 ==")
        walk(slide.shapes, parts)
        # 備忘稿常常才是真正的說明(投影片上只有幾個關鍵字)
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parts.append(f"== 第 {i} 頁的備忘稿 ==")
                parts.append(notes)
    return "\n".join(parts)


def _extract_docx(path: Path) -> str:
    """Word 轉文字:段落與**表格**,照文件裡的先後順序。

    ⚠️ python-docx 的 `doc.paragraphs` **不含表格裡的文字**,只讀它的話,
    欄位對照表這類文件(內容幾乎全在表格裡)只會抽到標題和幾句開場白 ——
    而 claude 不會說「我看不到表格」,它會拿剩下那點文字照樣回答。

    走 body 的子節點而不是先段落後表格:順序錯了的話「這張表在講哪一段」就斷了。
    """
    from docx import Document
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(str(path))
    parts: List[str] = []
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            text = Paragraph(child, doc).text.strip()
            if text:
                parts.append(text)
        elif child.tag == qn("w:tbl"):
            table = Table(child, doc)
            rows = [[cell.text for cell in row.cells] for row in table.rows]
            parts.extend(_table_lines(rows))
    return "\n".join(parts)


def _extract_xls(path: Path) -> str:
    """舊版 Excel(97-2003)。輸出格式跟 _extract_xlsx 一致,claude 那邊看起來一樣。

    xlrd 2.x **只**讀得了 .xls —— 它在 2.0 把 .xlsx 的支援拿掉了,所以這支不會
    跟 openpyxl 打架。日期在 .xls 裡是浮點數,轉回日期要靠活頁簿自己的 datemode。
    """
    import xlrd

    book = xlrd.open_workbook(str(path))
    parts = []
    for sheet in book.sheets():
        parts.append(f"== 工作表:{sheet.name} ==")
        for r in range(sheet.nrows):
            cells = []
            for cell in sheet.row(r):
                if cell.ctype == xlrd.XL_CELL_DATE:
                    try:
                        cells.append(
                            datetime.datetime(
                                *xlrd.xldate_as_tuple(cell.value, book.datemode)
                            ).isoformat(sep=" ")
                        )
                        continue
                    except (ValueError, OverflowError, xlrd.XLDateError):
                        pass  # 壞掉的日期就照原始數值印,不要整份抽不出來
                if cell.ctype == xlrd.XL_CELL_EMPTY:
                    cells.append("")
                elif cell.ctype == xlrd.XL_CELL_NUMBER and float(cell.value).is_integer():
                    # xlrd 的數字一律是 float,不處理的話每個整數都會多一個 ".0"
                    cells.append(str(int(cell.value)))
                else:
                    cells.append(str(cell.value))
            if any(c for c in cells):
                parts.append("\t".join(cells))
    return "\n".join(parts)


def _extract_xlsx(path: Path) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(str(path), read_only=True, data_only=True)
    parts = []
    for ws in wb.worksheets:
        parts.append(f"== 工作表:{ws.title} ==")
        for row in ws.iter_rows(values_only=True):
            if any(cell is not None for cell in row):
                parts.append("\t".join("" if cell is None else str(cell) for cell in row))
    return "\n".join(parts)
