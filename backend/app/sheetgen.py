"""把 claude 寫出來的 .csv 另存一份 Excel(.xlsx)。

**為什麼這件事要由後端做**:`.xlsx` 是一包 zip,非跑程式產不出來,而網頁的一般帳號
只有 `Read,Grep,Glob`,沒有 Bash。純文字的 `.csv` 它本來就寫得出來
(`agent.py` 無條件疊了 Write 到 `outputs/{session_id}/`),所以這裡只補最後一哩:
收尾時把每個 `.csv` 轉成同名 `.xlsx`。走的是跟
`agent.apply_copy_manifest` 完全一樣的套路 —— claude 只負責產資料,產檔在後端。

**兩份都留**:.csv 不刪。答案的文字裡常會提到自己寫了哪個檔,刪掉的話使用者
在下載清單裡找不到那個名字;而且要貼進別的系統時 .csv 反而比較好用。
留下來的 .csv 用 Excel 打開一樣會跑公式,所以轉完之後把它的公式開頭加前綴(`neutralize_csv`)。

⚠️ **只讀資料,不執行任何東西。** 跟 slidegen 同一條界線:claude 產出的永遠是
**資料**,版面與型態由這裡決定。這一階段連 manifest 都不用,claude 什麼都不必學。

排版只做三件不會猜錯的事:標題列粗體、凍結標題列、欄寬照內容。**刻意不猜更多** ——
多工作表、數字格式、日期型態都得由模型指定才不會做錯,那需要另一份 manifest。
"""

import codecs
import csv
import io
import logging
import re
from pathlib import Path
from typing import List, Optional, Union

from . import config, outputs_common

logger = logging.getLogger(__name__)

# Excel 自己的上限。超過的話 openpyxl 會在存檔時才爆,訊息也看不懂,
# 所以在這裡先截斷並講清楚截了多少。
MAX_ROWS = 1_048_576
MAX_COLS = 16_384

# 轉檔前先看檔案大小。收尾是關鍵路徑(使用者正盯著畫面等 done),
# 一份幾百 MB 的 csv 整包讀進記憶體再轉,會把整題卡在最後一步。
MAX_BYTES = 20 * 1024 * 1024

# 欄寬掃描的列數上限:再多也只是微調寬度,不值得為它多走十萬列。
WIDTH_SAMPLE_ROWS = 200
MAX_COL_WIDTH = 60

# 讀檔的編碼順序。cp950 一定要在:繁體中文 Windows 上 Excel 另存的 csv
# 多半是 Big5,用 utf-8 硬讀會整份變亂碼。
# 最後一道 utf-8/replace 不會失敗,所以這個函式不會因為編碼而丟例外。
_ENCODINGS = ("utf-8-sig", "cp950", "utf-8")

# 純數字才轉成數值。**開頭是 0 的不轉** —— 料號、廠商代號、部門代號這些常是
# 前面補零的字串,轉成數字等於把 "0012345" 變成 12345,拿去對帳就對不起來了。
# 千分位、貨幣符號、百分比一律留原樣:猜錯的代價比不猜大。
_NUMBER_RE = re.compile(r"-?(?:0|[1-9][0-9]*)(?:\.[0-9]+)?$")

# 工作表名稱不能有這些字,而且不能超過 31 個字(Excel 的規定,違反會開不起來)
_BAD_SHEET_CHARS = set("[]:*?/\\")


def _read_text(path: Path) -> str:
    """把 csv 讀成文字。編碼用猜的,但最後一道保證不會失敗。"""
    return _decode(path.read_bytes())[0]


def _decode(raw: bytes) -> "tuple[str, str]":
    """(文字, 寫回去要用的編碼)。utf-8-sig 也解得開沒有 BOM 的 utf-8,所以寫回去的編碼看原本有沒有 BOM。"""
    for encoding in _ENCODINGS:
        try:
            text = raw.decode(encoding)
        except UnicodeDecodeError:
            continue
        if encoding == "utf-8-sig" and not raw.startswith(codecs.BOM_UTF8):
            encoding = "utf-8"
        return text, encoding
    return raw.decode("utf-8", errors="replace"), "utf-8"


# Excel / LibreOffice 打開 csv 時,這幾個字開頭的格子會被當成公式(或 DDE)執行(OWASP CSV Injection)。
_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


#: 純數字:正負號、千分位、小數、百分比(-123、-1,234.5、+5%、-0012)。Excel 只會把它當數字,不會執行任何東西。
#: 跟 `_NUMBER_RE` 不是同一件事:那支決定 **xlsx 要不要轉成數值**(刻意保守,補零與千分位都留文字);
#: 這支只回答「這格有沒有危險」。
_PLAIN_NUMBER_RE = re.compile(r"[-+]?(?:[0-9]{1,3}(?:,[0-9]{3})+|[0-9]+)(?:\.[0-9]+)?%?$")


def _needs_guard(cell: str) -> bool:
    if len(cell) < 2 or not cell.startswith(_FORMULA_PREFIXES):
        return False  # 單獨一個「-」是報表常見的空值,Excel 也當文字
    # 負數(-123、-1,234.5)照原樣,不然整欄金額變文字、加總全錯
    return not _PLAIN_NUMBER_RE.match(cell)


def neutralize_csv(path: Path) -> bool:
    """留下來給人下載的 .csv:公式開頭的格子前面加一個 `'`(Excel 打開顯示成文字、不執行)。回傳有沒有改。

    **只在真的有那種格子時才重寫**,其他格子、分隔符號、編碼、換行都照原本的(csv 常要貼進別的系統)。
    xlsx 不用這一步:openpyxl 只把「=」開頭當公式,`csv_to_xlsx` 已經強制成文字。
    """
    raw = path.read_bytes()
    text, encoding = _decode(raw)
    delimiter = _sniff_delimiter(text)
    rows = list(csv.reader(io.StringIO(text, newline=""), delimiter=delimiter))
    changed = False
    for row in rows:
        for i, cell in enumerate(row):
            if _needs_guard(cell):
                row[i] = "'" + cell
                changed = True
    if not changed:
        return False
    out = io.StringIO()
    newline = "\r\n" if "\r\n" in text else "\n"
    csv.writer(out, delimiter=delimiter, lineterminator=newline).writerows(rows)
    path.write_bytes(out.getvalue().encode(encoding, errors="replace"))
    return True


def _sniff_delimiter(text: str) -> str:
    """猜分隔符號。

    不用 csv.Sniffer:它對「只有一欄」或「欄位裡有標點」的檔案很容易猜歪,
    而且猜不出來時是丟例外。這裡只在三個候選裡數次數,猜不出來就用逗號 ——
    猜錯的結果是「整列擠在 A 欄」,看得出來也救得回來。
    """
    lines = [line for line in text.splitlines()[:5] if line.strip()]
    if not lines:
        return ","
    best, best_count = ",", 0
    for candidate in (",", "\t", ";"):
        count = min(line.count(candidate) for line in lines)
        if count > best_count:
            best, best_count = candidate, count
    return best


def _cell_value(value: str) -> Union[str, int, float]:
    """該是數字的才變數字,其餘一律原樣留成文字。"""
    text = value.strip()
    if not text or not _NUMBER_RE.match(text):
        return value
    # 超過 15 位數的整數在 Excel 裡會被浮點數吃掉尾數(統一編號、序號踩得到),
    # 寧可留成文字。
    if len(text.lstrip("-").replace(".", "")) > 15:
        return value
    return float(text) if "." in text else int(text)


def _display_width(text: str) -> int:
    """估算顯示寬度。中文字算兩格,不然整欄中文的欄寬會只有一半。"""
    return sum(2 if ord(c) > 0x2E7F else 1 for c in text)


def _sheet_title(stem: str) -> str:
    cleaned = "".join(c for c in stem if c not in _BAD_SHEET_CHARS).strip()
    return cleaned[:31] or "工作表1"


def csv_to_xlsx(src: Path, dst: Path) -> List[str]:
    """把一份 csv 轉成 xlsx。回傳要顯示給使用者的訊息(正常是空的)。"""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    notices: List[str] = []
    text = _read_text(src)
    rows = list(csv.reader(io.StringIO(text), delimiter=_sniff_delimiter(text)))
    if not any(any(cell.strip() for cell in row) for row in rows):
        return []

    if len(rows) > MAX_ROWS:
        notices.append(
            "「%s」有 %s 列,超過 Excel 的上限,%s 只放得下前 %s 列"
            % (src.name, len(rows), dst.name, MAX_ROWS)
        )
        rows = rows[:MAX_ROWS]

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = _sheet_title(src.stem)

    truncated_cols = False
    widths: List[int] = []
    for r, row in enumerate(rows, start=1):
        if len(row) > MAX_COLS:
            truncated_cols = True
            row = row[:MAX_COLS]
        for c, raw in enumerate(row, start=1):
            cell = sheet.cell(row=r, column=c, value=_cell_value(raw))
            # 「=」開頭的字串 openpyxl 會當成公式寫進去,打開就是 #NAME? 或更糟
            # (別人的試算表被塞進一條會跑的公式)。強制標成文字。
            if isinstance(cell.value, str) and cell.value.startswith("="):
                cell.data_type = "s"
        if r == 1:
            for cell in sheet[1]:
                cell.font = Font(bold=True)
        if r <= WIDTH_SAMPLE_ROWS:
            for c, raw in enumerate(row):
                width = min(_display_width(raw.strip()) + 2, MAX_COL_WIDTH)
                if c < len(widths):
                    widths[c] = max(widths[c], width)
                else:
                    widths.append(width)

    if truncated_cols:
        notices.append(
            "「%s」有超過 %s 欄,超過的部分沒有放進 %s" % (src.name, MAX_COLS, dst.name)
        )

    for c, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(c)].width = max(width, 6)
    # 只有一列的時候不凍結:凍住唯一一列等於什麼都看不到
    if len(rows) > 1:
        sheet.freeze_panes = "A2"

    workbook.save(str(dst))
    return notices


def apply_csv_exports(session_id: str, knowledge_base: Optional[dict] = None) -> List[str]:
    """把這一題產出的每個 .csv 都另存一份同名 .xlsx,回傳要顯示給使用者的訊息。

    跟其他收尾後處理一樣,由 `routers/ask.py` 的 `_finalize` 呼叫,
    而且**一定要排在 `list_output_files` 前面** —— 不然轉出來的 .xlsx
    不會出現在下載清單裡。

    任何一步失敗都只回訊息、不丟例外:這是收尾階段,答案已經算完了,
    不該因為轉檔失敗把整題變成失敗。

    `knowledge_base` 這一支用不到,但收尾的後處理**簽名一律相同**
    (`apply_copy_manifest` 真的會用到它,要靠它驗證來源路徑),`_finalize` 才能拿一個 tuple 逐支呼叫。
    """
    session_dir = config.OUTPUT_DIR / session_id
    if not session_dir.is_dir():
        return []

    notices: List[str] = []
    sources = [p for p in sorted(session_dir.glob("*.csv")) if p.is_file()]
    for index, src in enumerate(sources):
        # 撞名就加序號,**不要蓋也不要跳過**:蓋掉會毀了 copy manifest 剛複製進來的
        # 正本;跳過則是使用者要了 Excel 卻什麼都沒拿到,而畫面上一句話都沒有。
        # 加序號兩份都留得住,而且下載清單上看得出來。
        dst = outputs_common.unique_path(session_dir, src.stem + ".xlsx")
        try:
            if src.stat().st_size > MAX_BYTES:
                notices.append(
                    "「%s」太大(超過 %d MB),沒有轉成 Excel,請直接下載 csv。"
                    % (src.name, MAX_BYTES // (1024 * 1024))
                )
                continue
            notices.extend(csv_to_xlsx(src, dst))
        except ImportError:
            # 一份都轉不了,但後面還沒輪到的 csv 一樣留在下載清單上,一樣要擋公式(這一份由 finally 擋)
            for rest in sources[index + 1:]:
                _guard(rest, session_id)
            return notices + outputs_common.missing_package("openpyxl", "Excel 檔")
        except Exception as exc:  # noqa: BLE001 —— 收尾階段,不能讓它把整題弄成失敗
            logger.exception("csv 轉 Excel 失敗 session=%s file=%s", session_id, src.name)
            notices.append(
                "「%s」轉成 Excel 失敗:%s: %s(csv 本身還是可以下載)"
                % (src.name, type(exc).__name__, exc)
            )
            outputs_common.discard_partial(dst)
        finally:
            # 轉完才改:xlsx 要拿原值(加了前綴的話 Excel 裡會看到那個 `'`)。太大沒轉的那份一樣要擋
            _guard(src, session_id)

    return notices


def _guard(src: Path, session_id: str) -> None:
    try:
        neutralize_csv(src)
    except Exception:  # noqa: BLE001 —— 收尾階段
        logger.exception("csv 公式防護失敗 session=%s file=%s", session_id, src.name)
