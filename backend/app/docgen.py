"""把 claude 寫出來的 Markdown 轉成 Word(.docx)。

跟 `pdfgen` 同一個理由與同一套做法:Word 是二進位,Write 工具寫不出來,而網頁的一般帳號
沒有 Bash 跑不了 python-docx。所以 claude 只負責寫一份 `.md`,再用 `__docx__.txt`
點名哪幾份要轉,真正產檔在這裡。

**為什麼吃 Markdown**:「請給我一份 Word」要的是一般文件,內容長什麼樣每次都不同。
Markdown 是模型本來就在寫的格式,不用教也不會寫錯;同一份 .md 也能同時轉 PDF。

**中文字型**:Word 不內嵌字型,東亞文字指定「微軟正黑體」(Windows 內建),
開檔的電腦沒有這個字型時 Word 會自己替換。⚠️ 只設 `font.name` 不夠 —— 那只改西文,
中文要另外寫 `w:eastAsia`;內建標題樣式還綁著佈景主題字型(`w:asciiTheme` 等),
不拿掉的話設了也不生效(見 `_set_fonts`)。

Markdown 怎麼切塊、清單怎麼跑跟 `pdfgen` 共用一份(`markdown_doc`),這裡只管畫成 Word;
另外多認 `>` 引用(套 Word 內建的 Quote 樣式)。行內只認 `**粗體**` 與 `` `等寬` ``,
**其餘語法一律照原字印出來** —— 看得懂比「安靜地漏掉」好。

編號清單照原本的數字印成「1. 文字」,不用 Word 的自動編號:python-docx 的
「List Number」整份文件共用一組編號,第二個清單會從上一個的結尾接著數。
"""
import logging
import re
from pathlib import Path
from typing import List, Optional, Sequence

from . import markdown_doc as md

logger = logging.getLogger(__name__)

# claude 用 Write 寫出來的「這幾份要轉 Word」清單。跟其他 manifest 一樣,
# 處理完就刪 —— 它是給後端看的中間產物,不該出現在使用者的下載清單裡。
MANIFEST_NAME = "__docx__.txt"

CJK_FONT = "Microsoft JhengHei"
LATIN_FONT = "Calibri"
MONO_FONT = "Consolas"

_INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")

_THEME_FONT_ATTRS = ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme")


def _set_fonts(element, latin: str, east_asian: str) -> None:
    """設一個樣式或一段文字的西文與東亞字型(`element` 是 style.element 或 run._element)。"""
    from docx.oxml.ns import qn

    rfonts = element.get_or_add_rPr().get_or_add_rFonts()
    for attr in _THEME_FONT_ATTRS:
        rfonts.attrib.pop(qn(attr), None)
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:eastAsia"), east_asian)


def _add_inline(paragraph, text: str) -> None:
    """把一行字(含 **粗體** 與 `等寬`)加進段落。"""
    for piece in _INLINE.split(text):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**") and len(piece) > 4:
            paragraph.add_run(piece[2:-2]).bold = True
        elif piece.startswith("`") and piece.endswith("`") and len(piece) > 2:
            run = paragraph.add_run(piece[1:-1])
            _set_fonts(run._element, MONO_FONT, CJK_FONT)
        else:
            paragraph.add_run(piece)


def _add_table(doc, rows: Sequence[Sequence[str]]) -> None:
    # 模型寫的表格少一根 | 是常有的事:補成一樣長,不要整份產不出來
    width = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=width)
    table.style = "Table Grid"
    for r, row in enumerate(rows):
        for c in range(width):
            cell = table.cell(r, c)
            paragraph = cell.paragraphs[0]
            _add_inline(paragraph, row[c] if c < len(row) else "")
            if r == 0:
                for run in paragraph.runs:
                    run.bold = True


def _add_rule(doc) -> None:
    """分隔線:一個只有下框線的空段落(Word 沒有「水平線」這種物件)。"""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    paragraph = doc.add_paragraph()
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for key, value in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "CCCCCC")):
        bottom.set(qn(key), value)
    border.append(bottom)
    paragraph._p.get_or_add_pPr().append(border)


def _add_code(doc, lines: Sequence[str]) -> None:
    """程式碼要逐行保留:同一段落裡用換行隔開,縮排的空白 python-docx 會原樣寫進去。"""
    code = doc.add_paragraph()
    for n, text in enumerate(lines):
        if n:
            code.add_run().add_break()
        _set_fonts(code.add_run(text)._element, MONO_FONT, CJK_FONT)


def _add_block(doc, block: md.Block) -> None:
    if isinstance(block, md.Paragraph):
        _add_inline(doc.add_paragraph(), block.text)
    elif isinstance(block, md.Heading):
        _add_inline(doc.add_heading("", level=block.level), block.text)
    elif isinstance(block, md.Code):
        _add_code(doc, block.lines)
    elif isinstance(block, md.Table):
        _add_table(doc, block.rows)
    elif isinstance(block, md.Quote):
        # 模型常用 `> 適用範圍說明` 這種開頭,照原字印的話 Word 裡會留一個 `>`;
        # 套 Word 內建的「Quote」樣式
        _add_inline(doc.add_paragraph(style="Quote"), block.text)
    elif isinstance(block, md.Bullets):
        for item in block.items:
            _add_inline(doc.add_paragraph(style="List Bullet"), item)
    elif isinstance(block, md.Numbered):
        _add_inline(doc.add_paragraph(), "%s. %s" % (block.number, block.text))
    elif isinstance(block, md.Rule):
        _add_rule(doc)
    elif isinstance(block, md.PageBreak):
        doc.add_page_break()


def render(markdown: str, out_path: Path, title: Optional[str] = None) -> List[str]:
    """把 Markdown 寫成 .docx,回傳要顯示給使用者的訊息。失敗會丟例外,由呼叫端收。"""
    from docx import Document

    doc = Document()
    doc.core_properties.title = title or out_path.stem
    doc.core_properties.author = "AskME"
    for name in ["Normal"] + ["Heading %d" % n for n in range(1, 7)]:
        _set_fonts(doc.styles[name].element, LATIN_FONT, CJK_FONT)

    blocks = md.parse(markdown, quotes=True)
    for block in blocks:
        _add_block(doc, block)
    # 分隔線與分頁不算內容:只有這兩種的話,開檔看到的是一份空白文件
    if all(isinstance(b, (md.Rule, md.PageBreak)) for b in blocks):
        doc.add_paragraph("(這份文件沒有內容)")
    doc.save(str(out_path))
    return []


def apply_docx_manifest(session_id: str, knowledge_base: Optional[dict] = None) -> List[str]:
    """把 `__docx__.txt` 點名的 Markdown 轉成 Word,回傳要顯示給使用者的訊息。

    由 `routers/ask.py` 的 `_finalize` 呼叫,**一定要排在 `list_output_files` 前面**
    —— 不然產出來的 .docx 不會出現在下載清單裡。失敗只回訊息、不丟例外
    (見 `markdown_doc.apply_manifest`)。
    """
    return md.apply_manifest(
        session_id,
        manifest_name=MANIFEST_NAME,
        label="Word",
        suffix=".docx",
        package="python-docx",
        render=render,
    )
